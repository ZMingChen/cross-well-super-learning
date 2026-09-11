import json
import os
import re
from pathlib import Path

import numpy as np
import pandas as pd


def norm_well_name(name):
    value = str(name).strip()
    value = re.sub(r"\.(las|txt)$", "", value, flags=re.I)
    return value.upper()


def workbook_rows(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def read_code_map(path):
    mapping = {}
    for row in workbook_rows(path)[1:]:
        if not row or row[0] is None:
            continue
        try:
            code = int(float(row[0]))
        except Exception:
            continue
        mapping[code] = str(row[1]).strip() if len(row) > 1 and row[1] is not None else str(code)
    return mapping


def read_intervals(path):
    records = []
    for row in workbook_rows(path)[1:]:
        if not row or row[0] is None:
            continue
        try:
            top = float(row[1])
            bottom = float(row[2])
            code = int(float(row[3]))
        except Exception:
            continue
        if bottom <= top:
            continue
        records.append({"label_well": norm_well_name(row[0]), "top": top, "bottom": bottom, "human_code": code})
    return pd.DataFrame(records).sort_values(["label_well", "top", "bottom"]).reset_index(drop=True)


def match_label_well(well, label_wells):
    if well in label_wells:
        return well
    candidates = []
    if well.endswith("X"):
        candidates.append(well[:-1] + "H")
    if well.endswith("H"):
        candidates.append(well[:-1] + "X")
    candidates += [well.replace("-2X", "-2H"), well.replace("-3X", "-3H")]
    for candidate in candidates:
        if candidate in label_wells:
            return candidate
    return None


def assign_human_codes(labels, intervals):
    label_wells = set(intervals["label_well"].unique())
    by_well = {w: p.sort_values("top") for w, p in intervals.groupby("label_well")}
    parts = []
    for well, part in labels.groupby("well", sort=True):
        matched = match_label_well(well, label_wells)
        out = part.copy()
        out["matched_label_well"] = matched if matched else ""
        out["human_code"] = np.nan
        if matched is None:
            parts.append(out)
            continue
        local = by_well[matched]
        depths = out["DEPTH"].to_numpy()
        tops = local["top"].to_numpy()
        bottoms = local["bottom"].to_numpy()
        codes = local["human_code"].to_numpy()
        idx = np.searchsorted(tops, depths, side="right") - 1
        assigned = np.full(len(out), np.nan)
        mask = (idx >= 0) & (idx < len(tops))
        pos = np.where(mask)[0]
        valid_idx = idx[mask]
        inside = depths[mask] < bottoms[valid_idx]
        assigned[pos[inside]] = codes[valid_idx[inside]]
        out["human_code"] = assigned
        parts.append(out)
    return pd.concat(parts, ignore_index=True, sort=False)


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Attach manual interval labels to a feature table.")
    parser.add_argument("--labels", type=Path, required=True, help="Workbook: well, top, bottom, code in first four columns.")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = pd.read_csv(args.input, encoding="utf-8-sig")
    data["DEPTH"] = pd.to_numeric(data["DEPTH"], errors="raise")
    assigned = assign_human_codes(data, read_intervals(args.labels))
    assigned["final_code"] = assigned["human_code"]
    assigned["final_label_source"] = np.where(assigned["human_code"].notna(), "human", "unlabelled")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    assigned.to_csv(args.output, index=False, encoding="utf-8-sig")


if __name__ == "__main__":
    main()
